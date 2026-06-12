import json
import re
import unicodedata
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


WORKSPACE = Path(r"C:\Users\tonym\OneDrive\Documentos\PPGEF - DashBorad")
INPUT_XLSX = Path(
    r"G:\.shortcut-targets-by-id\18B_rphPE1Flt1tGbECpP70xTWYIcUsfW\AAA - PPGEF EFI\Docentes\AAA - Corpo Docente PPGEF 2.0.xlsb.xlsx"
)
DATA_JSON = WORKSPACE / "data" / "corpo_docente.json"
DATA_JS = WORKSPACE / "dashboard" / "corpo-docente-data.js"
REPORT_JSON = WORKSPACE / "analysis" / "corpo_docente_indicator_report.json"

YEAR_COLUMNS = [str(year) for year in range(2021, 2029)]
THRESHOLDS = {
    "permanentes_no_corpo": {"target": 0.70, "direction": "min"},
    "permanentes_internos": {"target": 0.80, "direction": "min"},
    "permanentes_exclusivos": {"target": 0.30, "direction": "min"},
    "permanentes_externos": {"target": 0.20, "direction": "max"},
    "permanentes_colaboradores_outros_ppgs": {"target": 0.40, "direction": "max"},
    "total_permanentes": {"target": 12, "direction": "min"},
}


def norm_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\xa0", " ").strip()
    text = re.sub(r"\s+", " ", text)
    return text


def fold_text(value: Any) -> str:
    text = norm_text(value).lower()
    text = "".join(
        char for char in unicodedata.normalize("NFKD", text) if not unicodedata.combining(char)
    )
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def token_set(value: Any) -> set[str]:
    return {token for token in fold_text(value).split() if token}


def bool_or_none(value: Any) -> bool | None:
    if value in (True, False):
        return value
    text = norm_text(value).lower()
    if text in {"true", "sim", "s"}:
        return True
    if text in {"false", "nao", "não", "n"}:
        return False
    return None


def numeric_or_none(value: Any) -> int | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return int(value)
    text = norm_text(value)
    if text in {"", "---", "???"}:
        return None
    try:
        return int(float(text.replace(",", ".")))
    except ValueError:
        return None


def parse_date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    text = norm_text(value)
    if not text:
        return ""
    for fmt in ("%Y-%m-%d %H:%M:%S", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return text


def evaluate_threshold(metric_key: str, value: float | None) -> dict[str, Any]:
    config = THRESHOLDS[metric_key]
    target = config["target"]
    if metric_key == "total_permanentes":
        fallback_min = 10
        if value is None:
            passed = False
            gap = None
            condition = "indefinido"
        elif value >= target:
            passed = True
            gap = value - target
            condition = "ok"
        elif value >= fallback_min:
            passed = False
            gap = value - target
            condition = "justificavel_regional"
        else:
            passed = False
            gap = value - target
            condition = "abaixo_do_minimo"
        return {
            "metric": metric_key,
            "target": target,
            "direction": config["direction"],
            "value": value,
            "passed": passed,
            "gap": gap,
            "fallback_min": fallback_min,
            "condition": condition,
        }

    if value is None:
        passed = False
        gap = None
    elif config["direction"] == "min":
        passed = value >= target
        gap = value - target
    else:
        passed = value <= target
        gap = target - value
    return {
        "metric": metric_key,
        "target": target,
        "direction": config["direction"],
        "value": value,
        "passed": passed,
        "gap": gap,
    }


def read_cadastro_sheet(workbook) -> dict[str, dict[str, Any]]:
    ws = workbook["Cadastro"]
    headers = [norm_text(ws.cell(6, col).value) for col in range(2, 14)]
    rows = {}
    for row_index in range(7, ws.max_row + 1):
        name = norm_text(ws.cell(row_index, 2).value)
        if not name:
            continue
        row = {
            headers[col - 2]: norm_text(ws.cell(row_index, col).value) for col in range(2, 14)
        }
        row["row_index"] = row_index
        row["name_key"] = fold_text(name)
        rows[row["name_key"]] = row
    return rows


def resolve_by_name_key(rows: dict[str, dict[str, Any]], raw_name: str) -> dict[str, Any] | None:
    key = fold_text(raw_name)
    exact = rows.get(key)
    if exact:
        return exact

    sought_tokens = token_set(raw_name)
    if len(sought_tokens) < 2:
        return None

    candidates = []
    for row in rows.values():
        candidate_tokens = token_set(row.get("Nome do Docente") or row.get("docente_lattes") or "")
        if sought_tokens.issubset(candidate_tokens):
            candidates.append(row)
    if len(candidates) == 1:
        return candidates[0]
    return None


def read_atuacao_sheet(workbook) -> list[dict[str, Any]]:
    ws = workbook["2025-2028"]
    headers = [norm_text(ws.cell(5, col).value) for col in range(2, 22)]
    rows = []
    for row_index in range(6, 37):
        name = norm_text(ws.cell(row_index, 2).value)
        if not name:
            continue
        raw = {headers[col - 2]: ws.cell(row_index, col).value for col in range(2, 22)}
        row = {
            "docente": norm_text(raw["Docente"]),
            "name_key": fold_text(raw["Docente"]),
            "ams": bool_or_none(raw["AMS"]),
            "dfe": bool_or_none(raw["DFE"]),
            "origem_matriz": norm_text(raw["Origem"]),
            "ppgs_como_permanente_2026": numeric_or_none(raw["# PPGs como P em 2026"]),
            "colabora_outros_ppgs": bool_or_none(raw["COL em outros PPGs"]),
            "status_quadri": norm_text(raw["Status Quadri"]),
            "anos": {year: norm_text(raw[year]) for year in YEAR_COLUMNS},
            "row_index": row_index,
        }
        rows.append(row)
    return rows


def read_lattes_sheet(workbook) -> dict[str, dict[str, Any]]:
    ws = workbook["Atualização Lattes"]
    rows = {}
    cutoff = datetime(2024, 12, 1).date()
    for row_index in range(4, ws.max_row + 1):
        name = norm_text(ws.cell(row_index, 2).value)
        if not name:
            continue
        update_date = parse_date(ws.cell(row_index, 3).value)
        parsed_date = None
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", update_date):
            parsed_date = datetime.strptime(update_date, "%Y-%m-%d").date()
        status_value = norm_text(ws.cell(row_index, 4).value)
        status = status_value if status_value in {"Ok", "Revisar"} else ""
        if not status and parsed_date:
            status = "Revisar" if parsed_date < cutoff else "Ok"
        rows[fold_text(name)] = {
            "docente_lattes": name,
            "data_atualizacao": update_date,
            "status_lattes": status or "Indefinido",
            "row_index": row_index,
        }
    return rows


def category_from_year_status(status: str) -> str:
    mapping = {
        "Per": "Permanente",
        "Col": "Colaborador",
        "Vis": "Visitante",
        "Temp": "Temporário",
    }
    return mapping.get(status, "")


def build_docente_records() -> dict[str, Any]:
    workbook = load_workbook(INPUT_XLSX, data_only=True, read_only=True)
    cadastro = read_cadastro_sheet(workbook)
    atuacao_rows = read_atuacao_sheet(workbook)
    lattes = read_lattes_sheet(workbook)

    docentes = []
    issues = []

    for row in atuacao_rows:
        cadastro_row = resolve_by_name_key(cadastro, row["docente"])
        lattes_row = resolve_by_name_key(lattes, row["docente"])
        docente_issues = []

        if cadastro_row and norm_text(cadastro_row["Nome do Docente"]) != row["docente"]:
            docente_issues.append("nome_divergente_entre_abas")
        if lattes_row and norm_text(lattes_row["docente_lattes"]) != row["docente"]:
            docente_issues.append("nome_divergente_com_lattes")
        if not cadastro_row:
            docente_issues.append("ausente_no_cadastro")
        if row["status_quadri"] == "A" and cadastro_row and cadastro_row["Status"] != "A":
            docente_issues.append("status_divergente_ativo_inativo")

        cadastro_categoria = cadastro_row["Categoria"] if cadastro_row else ""
        categoria_2026 = category_from_year_status(row["anos"]["2026"])
        if cadastro_categoria and categoria_2026 and cadastro_categoria not in {"---", categoria_2026}:
            docente_issues.append("categoria_divergente_cadastro_vs_2026")

        cadastro_origem = cadastro_row["Origem 2"] if cadastro_row else ""
        if cadastro_origem == "UFPE" and row["origem_matriz"] == "Externo":
            docente_issues.append("origem_divergente_interno_vs_externo")
        if cadastro_origem == "Externo" and row["origem_matriz"] == "Interno":
            docente_issues.append("origem_divergente_externo_vs_interno")

        if cadastro_row and any(cadastro_row.get(field) in {"---", "???"} for field in ("Categoria", "ORCID")):
            docente_issues.append("cadastro_incompleto")
        if lattes_row and lattes_row["status_lattes"] == "Revisar":
            docente_issues.append("lattes_desatualizado")
        if row["anos"]["2026"] == "Per" and row["origem_matriz"] == "Externo":
            docente_issues.append("permanente_externo_2026")
        if row["anos"]["2026"] == "Per" and (row["ppgs_como_permanente_2026"] or 0) > 1:
            docente_issues.append("nao_exclusivo_ao_programa_2026")
        if row["anos"]["2026"] == "Per" and row["colabora_outros_ppgs"] is True:
            docente_issues.append("colabora_em_outro_ppg_2026")

        docente = {
            "docente": row["docente"],
            "name_key": row["name_key"],
            "status_quadri": row["status_quadri"],
            "ams": row["ams"],
            "dfe": row["dfe"],
            "origem_matriz": row["origem_matriz"],
            "ppgs_como_permanente_2026": row["ppgs_como_permanente_2026"],
            "colabora_outros_ppgs": row["colabora_outros_ppgs"],
            "anos": row["anos"],
            "cadastro": cadastro_row or {},
            "lattes": lattes_row or {},
            "issues": docente_issues,
        }
        docentes.append(docente)

        for code in docente_issues:
            issues.append({"docente": row["docente"], "issue": code})

    by_year = {}
    for year in YEAR_COLUMNS:
        classified = [doc for doc in docentes if doc["anos"].get(year) in {"Per", "Col", "Vis", "Temp"}]
        permanentes = [doc for doc in classified if doc["anos"][year] == "Per"]
        internos = [doc for doc in permanentes if doc["origem_matriz"] == "Interno"]
        externos = [doc for doc in permanentes if doc["origem_matriz"] == "Externo"]
        exclusivos = [
            doc for doc in permanentes if (doc["ppgs_como_permanente_2026"] or 0) == 1
        ]
        colaboradores_outros = [
            doc for doc in permanentes if doc["colabora_outros_ppgs"] is True
        ]
        statuses = Counter(doc["anos"][year] for doc in classified)

        metric_values = {
            "permanentes_no_corpo": (len(permanentes) / len(classified)) if classified else None,
            "permanentes_internos": (len(internos) / len(permanentes)) if permanentes else None,
            "permanentes_exclusivos": (len(exclusivos) / len(permanentes)) if permanentes else None,
            "permanentes_externos": (len(externos) / len(permanentes)) if permanentes else None,
            "permanentes_colaboradores_outros_ppgs": (
                len(colaboradores_outros) / len(permanentes)
            )
            if permanentes
            else None,
            "total_permanentes": len(permanentes),
        }

        by_year[year] = {
            "classified_count": len(classified),
            "permanent_count": len(permanentes),
            "status_counts": dict(statuses),
            "metrics": {
                key: evaluate_threshold(key, value) for key, value in metric_values.items()
            },
            "docentes": [doc["docente"] for doc in classified],
        }

    summary = {
        "source": str(INPUT_XLSX),
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "default_year": "2026",
        "thresholds": THRESHOLDS,
        "docentes": docentes,
        "by_year": by_year,
        "issue_counts": dict(Counter(issue["issue"] for issue in issues)),
        "issues": issues,
    }
    return summary


def main():
    summary = build_docente_records()
    DATA_JSON.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    DATA_JS.write_text(
        "window.CORPO_DOCENTE_PPGEF = " + json.dumps(summary, ensure_ascii=False) + ";",
        encoding="utf-8",
    )
    REPORT_JSON.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"default_year": summary["default_year"], "issue_counts": summary["issue_counts"], "by_year": summary["by_year"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
