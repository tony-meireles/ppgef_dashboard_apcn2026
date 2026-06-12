import csv
import json
from pathlib import Path

from openpyxl import load_workbook


WORKSPACE = Path(r"C:\Users\tonym\OneDrive\Documentos\PPGEF - DashBorad")
INPUT_XLSX = Path(
    r"G:\.shortcut-targets-by-id\18B_rphPE1Flt1tGbECpP70xTWYIcUsfW\AAA - PPGEF EFI\Docentes\AAA - Corpo Docente PPGEF 2.0.xlsb.xlsx"
)
ANALYSIS_DIR = WORKSPACE / "analysis"


def norm(value):
    if value is None:
        return ""
    return str(value).replace("\xa0", " ").strip()


def first_nonempty_cells(row):
    values = []
    for index, cell in enumerate(row, start=1):
        text = norm(cell)
        if text:
            values.append({"col": index, "value": text})
    return values


def sheet_preview(ws, max_rows=15, max_cols=40):
    preview = []
    for row in ws.iter_rows(min_row=1, max_row=max_rows, max_col=max_cols, values_only=True):
        preview.append(first_nonempty_cells(row))
    return preview


def export_sheet_csv(ws, destination, max_cols=120):
    with destination.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        for row in ws.iter_rows(values_only=True, max_col=max_cols):
            writer.writerow([norm(cell) for cell in row])


def main():
    ANALYSIS_DIR.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(INPUT_XLSX, data_only=False, read_only=True)

    summary = {"source": str(INPUT_XLSX), "sheets": []}
    target_sheets = {"Cadastro", "2025-2028", "Simulação 3 LP", "Simulações", "xxx - Ajustes 2025"}

    for ws in workbook.worksheets:
        info = {
            "name": ws.title,
            "max_row": ws.max_row,
            "max_col": ws.max_column,
            "preview": sheet_preview(ws),
        }
        summary["sheets"].append(info)

        if ws.title in target_sheets:
            csv_name = f"corpo_docente_{ws.title.lower().replace(' ', '_').replace('-', '_')}.csv"
            export_sheet_csv(ws, ANALYSIS_DIR / csv_name)

    (ANALYSIS_DIR / "corpo_docente_workbook_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
