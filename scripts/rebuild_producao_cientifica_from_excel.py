import csv
import json
import re
import unicodedata
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


WORKSPACE = Path(r"C:\Users\tonym\OneDrive\Documentos\PPGEF - DashBorad")
INPUT_XLSX = Path(
    r"G:\.shortcut-targets-by-id\18B_rphPE1Flt1tGbECpP70xTWYIcUsfW\AAA - PPGEF EFI\Sucupira\014 - Pj APCN 03 - 2026\BDs Estratégicos\Produção Docente 2026_06.xlsx"
)
PRIOR_JSON = WORKSPACE / "data" / "producao_cientifica.json"
OUTPUT_CSV = WORKSPACE / "outputs" / "producao_cientifica_rebuild_staging.csv"
OUTPUT_JSON = WORKSPACE / "outputs" / "producao_cientifica_rebuild_staging.json"
OUTPUT_JS = WORKSPACE / "outputs" / "producao-cientifica-data.staging.js"
REPORT_JSON = WORKSPACE / "analysis" / "producao_cientifica_rebuild_staging_report.json"


def norm_text(value: Any) -> str:
    text = str(value or "").replace("\xa0", " ").replace("\u2009", " ").strip()
    text = re.sub(r"\s+", " ", text)
    return text


def fold_text(value: Any) -> str:
    text = norm_text(value).lower()
    text = "".join(
        ch for ch in unicodedata.normalize("NFKD", text) if not unicodedata.combining(ch)
    )
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def compact_text(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", fold_text(value))


def maybe_blank(value: Any) -> str:
    text = norm_text(value)
    return "" if text.lower() in {"", "nan", "none"} else text


def extract_doi(text: str) -> str:
    match = re.search(r"(10\.\d{4,9}/[-._;()/:A-Z0-9]+)", text, flags=re.I)
    return match.group(1).rstrip(" .;,)") if match else ""


def extract_year(text: str) -> str:
    matches = re.findall(r"\b(19|20)\d{2}\b", text)
    if not matches:
        return ""
    years = re.findall(r"\b(?:19|20)\d{2}\b", text)
    return years[-1] if years else ""


def extract_isbn_or_issn(text: str) -> str:
    match = re.search(r"\b(\d{4}-?\d{3}[\dXx])\b", text)
    return match.group(1) if match else ""


def split_reference(product: str) -> tuple[list[str], str]:
    text = norm_text(product)
    text = re.sub(r"\s*Citações:\s*\d+\|\d+\s*$", "", text, flags=re.I)
    match = re.match(r"^(?P<authors>.+?)(?:\s\.\s+|\.\.\s+)(?P<rest>.+)$", text)
    if match:
        authors_part = match.group("authors")
        rest = match.group("rest")
    else:
        parts = [piece.strip() for piece in re.split(r"\.\s+", text, maxsplit=1) if piece.strip()]
        if len(parts) == 2:
            authors_part, rest = parts[0], parts[1]
        else:
            authors_part, rest = text, ""
    authors = [norm_text(author).strip(" .;") for author in authors_part.split(";") if norm_text(author).strip(" .;")]
    return authors, norm_text(rest)


def infer_tipo_producao(tipo: str) -> str:
    tipo_f = fold_text(tipo)
    if "artigo" in tipo_f:
        return "artigo"
    if "capitulo" in tipo_f:
        return "capitulo"
    if "livro" in tipo_f:
        return "livro"
    return tipo_f or "artigo"


def parse_reference_fields(product: str, tipo: str) -> dict[str, str]:
    text = norm_text(product)
    authors, rest = split_reference(text)
    doi = extract_doi(text)
    year = extract_year(text)
    isbn_issn = extract_isbn_or_issn(text)
    title = ""
    vehicle = ""
    editora = ""

    if infer_tipo_producao(tipo) == "capitulo" and re.search(r"\bIn:", rest, flags=re.I):
        title = norm_text(re.split(r"\bIn:", rest, maxsplit=1, flags=re.I)[0]).strip(" .;:")
        after_in = norm_text(re.split(r"\bIn:", rest, maxsplit=1, flags=re.I)[1])
        vehicle_match = re.search(r"\(Org\.\)\.?\s*(.+?)\.\s*\d+\s*ed", after_in, flags=re.I)
        vehicle = norm_text(vehicle_match.group(1)).strip(" .;:") if vehicle_match else ""
        if not vehicle:
            after_org = re.split(r"\(Org\.\)\.?", after_in, maxsplit=1, flags=re.I)
            if len(after_org) == 2:
                vehicle = norm_text(re.split(r"\.\s+", after_org[1], maxsplit=1)[0]).strip(" .;:")
        if not vehicle:
            vehicle = norm_text(re.split(r"\.\s+", after_in, maxsplit=1)[0]).strip(" .;:")
        publisher_match = re.search(r"\d+\s*ed\.[^:]*:\s*([^,]+)", after_in, flags=re.I)
        editora = norm_text(publisher_match.group(1)).strip(" .;:") if publisher_match else ""
    elif infer_tipo_producao(tipo) == "livro":
        title_match = re.match(r"^(.*?)\.\s*\d+\s*\.?\s*ed\.", rest, flags=re.I)
        title = norm_text(title_match.group(1)).strip(" .;:") if title_match else ""
        if not title:
            title = norm_text(re.split(r"\.\s+", rest, maxsplit=1)[0]).strip(" .;:")
        publisher_match = re.search(r"\d+\s*\.?\s*ed\.\s*([^:]+:)?\s*([^,.]+)", rest, flags=re.I)
        editora = norm_text(publisher_match.group(2)).strip(" .;:") if publisher_match else ""
    else:
        parts = [norm_text(piece).strip(" .;:") for piece in re.split(r"\.\s+", rest) if norm_text(piece).strip(" .;:")]
        if parts:
            title = parts[0]
        if len(parts) > 1:
            vehicle = parts[1]
        if infer_tipo_producao(tipo) == "livro" and not vehicle:
            editora = title
        elif infer_tipo_producao(tipo) == "livro":
            editora = vehicle

    if not year:
        year = extract_year(rest)

    return {
        "autores": "; ".join(authors),
        "autor_principal": authors[0] if authors else "",
        "ano": year,
        "titulo": title,
        "revista_ou_veiculo": vehicle,
        "editora": editora,
        "isbn_issn": isbn_issn,
        "doi": doi,
        "produto_referencia": text,
    }


@dataclass
class PriorMatch:
    row: dict[str, Any] | None
    source: str


def load_prior_rows() -> list[dict[str, Any]]:
    if not PRIOR_JSON.exists():
        return []
    rows = json.loads(PRIOR_JSON.read_text(encoding="utf-8"))
    for row in rows:
        row["doi_key"] = norm_text(row.get("doi")).lower()
        row["title_key"] = fold_text(row.get("titulo"))
        row["title_compact"] = compact_text(row.get("titulo"))
        row["journal_key"] = fold_text(row.get("revista_ou_veiculo"))
        row["docente_key"] = fold_text(row.get("docente"))
        row["tipo_key"] = fold_text(row.get("tipo_producao"))
    return rows


def build_prior_indexes(prior_rows: list[dict[str, Any]]) -> dict[str, dict[str, dict[str, Any]]]:
    by_doi: dict[str, dict[str, Any]] = {}
    by_title: dict[str, dict[str, Any]] = {}
    by_title_docente: dict[str, dict[str, Any]] = {}
    for row in prior_rows:
      doi_key = row.get("doi_key", "")
      if doi_key and doi_key not in by_doi:
          by_doi[doi_key] = row
      title_key = row.get("title_key", "")
      if title_key and title_key not in by_title:
          by_title[title_key] = row
      title_docente_key = f"{title_key}||{row.get('docente_key', '')}||{row.get('tipo_key', '')}"
      if title_key and title_docente_key not in by_title_docente:
          by_title_docente[title_docente_key] = row
    return {
        "by_doi": by_doi,
        "by_title": by_title,
        "by_title_docente": by_title_docente,
    }


def find_prior_match(
    indexes: dict[str, dict[str, dict[str, Any]]],
    docente: str,
    tipo_producao: str,
    titulo: str,
    doi: str,
) -> PriorMatch:
    doi_key = norm_text(doi).lower()
    if doi_key and doi_key in indexes["by_doi"]:
        return PriorMatch(indexes["by_doi"][doi_key], "doi")

    title_key = fold_text(titulo)
    docente_key = fold_text(docente)
    tipo_key = fold_text(tipo_producao)
    title_docente_key = f"{title_key}||{docente_key}||{tipo_key}"
    if title_key and title_docente_key in indexes["by_title_docente"]:
        return PriorMatch(indexes["by_title_docente"][title_docente_key], "title_docente")
    if title_key and title_key in indexes["by_title"]:
        return PriorMatch(indexes["by_title"][title_key], "title")
    return PriorMatch(None, "")


def merge_with_prior(
    docente: str,
    tipo_producao: str,
    parsed: dict[str, str],
    prior_match: PriorMatch,
) -> dict[str, Any]:
    prior = prior_match.row or {}
    revista_ou_veiculo = parsed["revista_ou_veiculo"] or maybe_blank(prior.get("revista_ou_veiculo"))
    editora = parsed["editora"] or maybe_blank(prior.get("editora"))
    isbn_issn = parsed["isbn_issn"] or maybe_blank(prior.get("isbn_issn"))
    doi = parsed["doi"] or maybe_blank(prior.get("doi"))
    ano = parsed["ano"] or maybe_blank(prior.get("ano"))

    observations = [
        "origem=rebuild_excel_staging",
        f"match_prior={prior_match.source or 'none'}",
    ]
    if not parsed["autores"]:
        observations.append("autores_nao_extraidos")

    return {
        "docente": docente,
        "autores": parsed["autores"],
        "autor_principal": parsed["autor_principal"],
        "ano": ano,
        "titulo": parsed["titulo"],
        "tipo_producao": tipo_producao,
        "revista_ou_veiculo": revista_ou_veiculo,
        "editora": editora,
        "cidade": maybe_blank(prior.get("cidade")),
        "isbn_issn": isbn_issn,
        "doi": doi,
        "url": maybe_blank(prior.get("url")),
        "indexacao": maybe_blank(prior.get("indexacao")),
        "jif_jcr_2025": maybe_blank(prior.get("jif_jcr_2025")),
        "quartil_jcr_2025": maybe_blank(prior.get("quartil_jcr_2025")),
        "pontuacao": maybe_blank(prior.get("pontuacao")),
        "observacoes": " | ".join(observations),
        "produto_referencia": parsed["produto_referencia"],
    }


def build_dedupe_key(row: dict[str, Any]) -> str:
    title_key = compact_text(row.get("titulo"))
    doi_key = norm_text(row.get("doi")).lower()
    if doi_key:
        title_key = f"{title_key}||{doi_key}"
    return "||".join(
        [
            fold_text(row.get("docente")),
            fold_text(row.get("tipo_producao")),
            row.get("ano", ""),
            title_key,
        ]
    )


def row_quality(row: dict[str, Any]) -> tuple[int, int, int, int]:
    return (
        1 if row.get("doi") else 0,
        1 if row.get("indexacao") else 0,
        1 if row.get("autores") else 0,
        len(row.get("produto_referencia", "")),
    )


def rebuild_from_excel() -> dict[str, Any]:
    wb = load_workbook(INPUT_XLSX, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    prior_rows = load_prior_rows()
    prior_indexes = build_prior_indexes(prior_rows)

    rebuilt_rows: list[dict[str, Any]] = []
    dedupe_map: dict[str, dict[str, Any]] = {}
    dedupe_ignored: list[dict[str, Any]] = []
    match_counter = Counter()
    type_counter = Counter()

    for source_index, row in enumerate(ws.iter_rows(min_row=4, values_only=True), start=4):
        docente = maybe_blank(row[2] if len(row) > 2 else "")
        tipo_raw = maybe_blank(row[3] if len(row) > 3 else "")
        produto = maybe_blank(row[4] if len(row) > 4 else "")
        if not docente or not tipo_raw or not produto:
            continue

        tipo_producao = infer_tipo_producao(tipo_raw)
        type_counter[tipo_producao] += 1
        parsed = parse_reference_fields(produto, tipo_raw)
        prior_match = find_prior_match(
            prior_indexes,
            docente=docente,
            tipo_producao=tipo_producao,
            titulo=parsed["titulo"],
            doi=parsed["doi"],
        )
        merged = merge_with_prior(docente, tipo_producao, parsed, prior_match)
        merged["source_row"] = str(source_index)

        dedupe_key = build_dedupe_key(merged)
        if dedupe_key in dedupe_map:
            incumbent = dedupe_map[dedupe_key]
            if row_quality(merged) > row_quality(incumbent):
                dedupe_ignored.append(incumbent)
                dedupe_map[dedupe_key] = merged
            else:
                dedupe_ignored.append(merged)
            continue

        dedupe_map[dedupe_key] = merged
        match_counter[prior_match.source or "none"] += 1

    rebuilt_rows = list(dedupe_map.values())
    rebuilt_rows.sort(
        key=lambda item: (
            fold_text(item["docente"]),
            item["ano"],
            fold_text(item["titulo"]),
            fold_text(item["tipo_producao"]),
        )
    )

    for idx, item in enumerate(rebuilt_rows, start=1):
        item["id"] = str(idx)

    fieldnames = [
        "id",
        "docente",
        "autores",
        "autor_principal",
        "ano",
        "titulo",
        "tipo_producao",
        "revista_ou_veiculo",
        "editora",
        "cidade",
        "isbn_issn",
        "doi",
        "url",
        "indexacao",
        "jif_jcr_2025",
        "quartil_jcr_2025",
        "pontuacao",
        "observacoes",
        "produto_referencia",
        "source_row",
    ]

    with OUTPUT_CSV.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rebuilt_rows)

    OUTPUT_JSON.write_text(json.dumps(rebuilt_rows, ensure_ascii=False, indent=2), encoding="utf-8")
    OUTPUT_JS.write_text(
        "window.PRODUCAO_CIENTIFICA = " + json.dumps(rebuilt_rows, ensure_ascii=False) + ";",
        encoding="utf-8",
    )

    report = {
        "input_xlsx": str(INPUT_XLSX),
        "prior_rows": len(prior_rows),
        "source_rows": sum(type_counter.values()),
        "rebuilt_rows": len(rebuilt_rows),
        "dedupe_ignored_rows": len(dedupe_ignored),
        "types": dict(type_counter),
        "prior_match_sources": dict(match_counter),
        "rows_with_authors": sum(1 for item in rebuilt_rows if item["autores"]),
        "rows_with_doi": sum(1 for item in rebuilt_rows if item["doi"]),
        "rows_with_indexacao_reused": sum(1 for item in rebuilt_rows if item["indexacao"]),
        "rows_with_pontuacao_reused": sum(1 for item in rebuilt_rows if item["pontuacao"]),
        "sample_deduped_out": [
            {
                "docente": item["docente"],
                "ano": item["ano"],
                "titulo": item["titulo"],
                "source_row": item["source_row"],
            }
            for item in dedupe_ignored[:15]
        ],
        "outputs": {
            "csv": str(OUTPUT_CSV),
            "json": str(OUTPUT_JSON),
            "js": str(OUTPUT_JS),
        },
    }
    REPORT_JSON.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


def main() -> None:
    report = rebuild_from_excel()
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
