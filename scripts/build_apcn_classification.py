import json
import math
import re
import time
import unicodedata
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.parse import quote, quote_plus
from urllib.request import Request, urlopen

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


WORKSPACE = Path(r"C:\Users\tonym\OneDrive\Documentos\PPGEF - DashBorad")
INPUT_XLSX = Path(r"C:\Users\tonym\Downloads\Base_APCN_Pronta_Para_Indexacao.xlsx")
PRIOR_JSON = WORKSPACE / "data" / "producao_cientifica.json"
OUTPUT_XLSX = WORKSPACE / "outputs" / "classificacao_apcn_area21_producao_docente.xlsx"
CACHE_JSON = WORKSPACE / "data" / "journal_lookup_cache.json"
TODAY_STR = date.today().isoformat()

MANUAL = "Verificação manual necessária"


def norm_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\xa0", " ").replace("\u2009", " ").strip()
    return re.sub(r"\s+", " ", text)


def fold_text(value: Any) -> str:
    text = norm_text(value).lower()
    text = "".join(
        ch for ch in unicodedata.normalize("NFKD", text) if not unicodedata.combining(ch)
    )
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def maybe_blank(value: Any) -> str:
    text = norm_text(value)
    if text.lower() in {"", "nan", "none"}:
        return ""
    return text


def looks_numeric_only(value: str) -> bool:
    return bool(re.fullmatch(r"\d{4,}", value))


def extract_doi(text: str) -> str:
    match = re.search(r"(10\.\d{4,9}/[-._;()/:A-Z0-9]+)", text, flags=re.I)
    return match.group(1).rstrip(" .;,)") if match else ""


def extract_isbn_or_issn(text: str) -> str:
    match = re.search(r"\b(\d{4}-?\d{3}[\dXx])\b", text)
    return match.group(1) if match else ""


def infer_type(tipo: str, produto: str) -> str:
    tipo_f = fold_text(tipo)
    prod_f = fold_text(produto)
    if "capitulo" in tipo_f or "capitulos" in tipo_f:
        return "Capítulo"
    if "livro" in tipo_f:
        return "Livro"
    if "artigo" in tipo_f:
        return "Artigo em periódico"
    if " in: " in prod_f and "(org" in prod_f:
        return "Capítulo"
    if re.search(r"\b\d+\s*p\.?\b", prod_f) and not " in: " in prod_f and ". ed." in prod_f:
        return "Livro"
    if any(token in prod_f for token in [" v. ", " doi", " issn", "http://", "https://"]):
        return "Artigo em periódico"
    return "Outro"


def extract_title_from_produto(produto: str) -> str:
    text = norm_text(produto)
    if not text:
        return ""
    title_part = text
    if " . " in text:
        pieces = [piece.strip() for piece in text.split(" . ") if piece.strip()]
        if len(pieces) >= 2:
            title_part = pieces[1]
    if " In: " in title_part:
        title_part = title_part.split(" In: ", 1)[0].strip()
    for token in [" http://", " https://", " DOI:", " doi:", " ISSN", " Citações:"]:
        idx = title_part.find(token)
        if idx >= 0:
            title_part = title_part[:idx].strip(" .;:")
    return title_part.strip(" .;:")


@dataclass
class JournalLookup:
    journal_query: str
    issn_query: str
    scimago_title: str = ""
    scimago_link: str = ""
    sjr_scopus: str = "Não"
    nlm_title: str = ""
    nlm_link: str = ""
    pubmed: str = "Não"
    pubmed_issn: str = ""
    jcr_wos: str = "Não"
    scielo: str = "Não"
    google_scholar: str = MANUAL
    observation: str = ""

    def to_dict(self) -> dict[str, str]:
        return self.__dict__.copy()


def url_json(url: str) -> Any:
    req = Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urlopen(req, timeout=25) as response:
        return json.load(response)


def url_text(url: str, referer: str = "") -> str:
    headers = {"User-Agent": "Mozilla/5.0"}
    if referer:
        headers["Referer"] = referer
    req = Request(url, headers=headers)
    with urlopen(req, timeout=25) as response:
        return response.read().decode("utf-8", "ignore")


def load_prior_metadata() -> tuple[pd.DataFrame, dict[str, str], dict[str, str]]:
    if not PRIOR_JSON.exists():
        return pd.DataFrame(), {}, {}
    prior = pd.DataFrame(json.loads(PRIOR_JSON.read_text(encoding="utf-8")))
    prior["journal_key"] = prior["revista_ou_veiculo"].map(fold_text)
    prior["title_key"] = prior["titulo"].map(fold_text)
    prior["doi_key"] = prior["doi"].fillna("").map(lambda x: norm_text(x).lower())

    journal_issn: dict[str, str] = {}
    journal_doi_hint: dict[str, str] = {}
    for _, row in prior.iterrows():
        jk = row["journal_key"]
        if jk and not journal_issn.get(jk):
            issn = maybe_blank(row.get("isbn_issn"))
            if issn:
                journal_issn[jk] = issn
        if jk and not journal_doi_hint.get(jk):
            doi = maybe_blank(row.get("doi"))
            if doi:
                journal_doi_hint[jk] = doi
    return prior, journal_issn, journal_doi_hint


def find_prior_match(
    prior: pd.DataFrame, journal: str, title: str, doi: str
) -> dict[str, str] | None:
    if prior.empty:
        return None
    doi_key = norm_text(doi).lower()
    if doi_key:
        matched = prior[prior["doi_key"] == doi_key]
        if not matched.empty:
            return matched.iloc[0].to_dict()

    title_key = fold_text(title)
    if title_key:
        matched = prior[prior["title_key"] == title_key]
        if not matched.empty:
            return matched.iloc[0].to_dict()

    journal_key = fold_text(journal)
    if journal_key:
        matched = prior[prior["journal_key"] == journal_key]
        if not matched.empty:
            return matched.iloc[0].to_dict()
    return None


def lookup_scimago(title: str, issn: str) -> tuple[str, str, str]:
    queries = []
    if issn:
        queries.append((issn.replace("-", ""), "iss"))
    if title:
        queries.append((title, "jou"))

    query_fold = fold_text(title)
    for raw_query, tip in queries:
        url = f"https://www.scimagojr.com/journalsearch.php?q={quote_plus(raw_query)}&tip={tip}"
        html = url_text(url, referer="https://www.scimagojr.com/")
        match_iter = re.finditer(
            r'<a href="(journalsearch\.php\?q=(\d+)&tip=sid&clean=0)">\s*<div class="sresult_item"><span class="jrnlname">(.*?)</span>(.*?)<br\s*/?>(.*?)</div>',
            html,
            flags=re.I | re.S,
        )
        for match in match_iter:
            found_title = norm_text(re.sub(r"<.*?>", "", match.group(3)))
            found_fold = fold_text(found_title)
            if tip == "iss":
                return "Sim", f"https://www.scimagojr.com/{match.group(1)}", found_title
            if query_fold and (
                found_fold == query_fold
                or found_fold.startswith(query_fold)
                or query_fold.startswith(found_fold)
            ):
                return "Sim", f"https://www.scimagojr.com/{match.group(1)}", found_title
    return "Não", "", ""


def lookup_nlm(title: str, issn: str) -> tuple[str, str, str, str]:
    queries = []
    if issn:
        queries.append(("issn", f'"{issn}"[ISSN]'))
        if "-" in issn:
            queries.append(("issn", f'"{issn.replace("-", "")}"[ISSN]'))
    if title:
        queries.append(("title", f'"{title}"[Title]'))

    title_fold = fold_text(title)
    for mode, term in queries:
        search_url = (
            "https://eutils.ncbi.nlm.nih.gov/entrez/eutils/esearch.fcgi?db=nlmcatalog&term="
            + quote(term)
            + "&retmode=json&retmax=20"
        )
        payload = url_json(search_url)
        ids = payload.get("esearchresult", {}).get("idlist", [])
        if not ids:
            continue
        summary_url = (
            "https://eutils.ncbi.nlm.nih.gov/entrez/eutils/esummary.fcgi?db=nlmcatalog&id="
            + ",".join(ids[:20])
            + "&retmode=json"
        )
        summary = url_json(summary_url).get("result", {})
        for uid in ids[:20]:
            item = summary.get(uid, {})
            title_list = item.get("titlemainlist") or []
            found_title = norm_text(title_list[0].get("title", "")) if title_list else ""
            found_fold = fold_text(found_title.rstrip("."))
            exact = mode == "issn" or (title_fold and found_fold == title_fold)
            if not exact:
                continue
            status = item.get("currentindexingstatus", "")
            found_issn = ""
            for issn_item in item.get("issnlist") or []:
                if issn_item.get("validyn") == "Y":
                    found_issn = issn_item.get("issn", "")
                    break
            link = f"https://www.ncbi.nlm.nih.gov/nlmcatalog/{uid}"
            if status == "Y":
                return "Sim", link, found_title.rstrip("."), found_issn
            return "Não", link, found_title.rstrip("."), found_issn
    return "Não", "", "", ""


def build_journal_cache(unique_journals: pd.DataFrame) -> dict[str, dict[str, str]]:
    existing: dict[str, dict[str, str]] = {}
    if CACHE_JSON.exists():
        existing = json.loads(CACHE_JSON.read_text(encoding="utf-8"))

    updated = dict(existing)
    for _, row in unique_journals.iterrows():
        key = row["journal_key"]
        if key in updated:
            continue

        lookup = JournalLookup(
            journal_query=row["Periodico_extraido"],
            issn_query=row["ISSN_identificado"],
        )
        observations: list[str] = []

        try:
            sjr, scimago_link, scimago_title = lookup_scimago(
                row["Periodico_extraido"], row["ISSN_identificado"]
            )
            lookup.sjr_scopus = sjr
            lookup.scimago_link = scimago_link
            lookup.scimago_title = scimago_title
        except (HTTPError, URLError, TimeoutError) as exc:
            lookup.sjr_scopus = MANUAL
            observations.append(f"SCImago indisponível: {exc}")

        try:
            pubmed, nlm_link, nlm_title, nlm_issn = lookup_nlm(
                row["Periodico_extraido"], row["ISSN_identificado"]
            )
            lookup.pubmed = pubmed
            lookup.nlm_link = nlm_link
            lookup.nlm_title = nlm_title
            lookup.pubmed_issn = nlm_issn
        except (HTTPError, URLError, TimeoutError) as exc:
            lookup.pubmed = MANUAL
            observations.append(f"NLM Catalog indisponível: {exc}")

        if row.get("Prior_WoS") == "Sim":
            lookup.jcr_wos = MANUAL
            observations.append("Base anterior indicava WoS/JCR, sem reconfirmação automática em Clarivate.")

        if lookup.sjr_scopus == "Não" and lookup.pubmed == "Não":
            observations.append("Sem confirmação automática em SCImago ou NLM; SciELO/Google Scholar dependem de verificação manual.")

        lookup.observation = " ".join(observations).strip()
        updated[key] = lookup.to_dict()
        time.sleep(0.25)

    CACHE_JSON.write_text(json.dumps(updated, ensure_ascii=False, indent=2), encoding="utf-8")
    return updated


def classify_articles(df: pd.DataFrame) -> pd.DataFrame:
    prior_df, journal_issn, _ = load_prior_metadata()

    df["Produto_norm"] = df["Produto"].map(norm_text)
    df["Tipo_APCN"] = [
        infer_type(tipo, produto) for tipo, produto in zip(df["Tipo"].fillna(""), df["Produto_norm"])
    ]
    df["Periodico_extraido"] = df["Periódico"].fillna("").map(maybe_blank)
    df.loc[df["Periodico_extraido"].map(looks_numeric_only), "Periodico_extraido"] = ""
    df["DOI_identificado"] = df["Produto_norm"].map(extract_doi)
    df["ISSN_identificado"] = ""
    df["Categoria_livro_capitulo"] = ""

    titles = []
    prior_wos_flags = []
    type_inconsistency = []
    for idx, row in df.iterrows():
        produto = row["Produto_norm"]
        periodico = row["Periodico_extraido"]
        title = extract_title_from_produto(produto)
        titles.append(title)

        prior_match = find_prior_match(prior_df, periodico, title, row["DOI_identificado"])
        if prior_match:
            if not periodico:
                candidate_journal = maybe_blank(prior_match.get("revista_ou_veiculo"))
                if candidate_journal:
                    df.at[idx, "Periodico_extraido"] = candidate_journal
                    periodico = candidate_journal
            if not df.at[idx, "DOI_identificado"]:
                doi = maybe_blank(prior_match.get("doi"))
                if doi:
                    df.at[idx, "DOI_identificado"] = doi
            issn = maybe_blank(prior_match.get("isbn_issn"))
            if issn:
                df.at[idx, "ISSN_identificado"] = issn
            prior_wos_flags.append(
                "Sim" if maybe_blank(prior_match.get("indexacao")) == "Web of Science (JCR)" else "Não"
            )
        else:
            prior_wos_flags.append("Não")
            journal_key = fold_text(periodico)
            if journal_key and journal_key in journal_issn:
                df.at[idx, "ISSN_identificado"] = journal_issn[journal_key]

        if not df.at[idx, "ISSN_identificado"]:
            extracted_issn = extract_isbn_or_issn(produto)
            if extracted_issn:
                df.at[idx, "ISSN_identificado"] = extracted_issn

        original_type = fold_text(row["Tipo"])
        inferred = df.at[idx, "Tipo_APCN"]
        inconsistent = "Não"
        if original_type and inferred == "Capítulo" and "cap" not in original_type:
            inconsistent = "Sim"
        if original_type and inferred == "Livro" and "livr" not in original_type:
            inconsistent = "Sim"
        type_inconsistency.append(inconsistent)

    df["Titulo_extraido"] = titles
    df["Prior_WoS"] = prior_wos_flags
    df["Possivel_inconsistencia_tipo"] = type_inconsistency

    article_rows = df["Tipo_APCN"] == "Artigo em periódico"
    unique_journals = (
        df.loc[article_rows & df["Periodico_extraido"].ne(""), ["Periodico_extraido", "ISSN_identificado", "Prior_WoS"]]
        .drop_duplicates()
        .assign(journal_key=lambda x: x["Periodico_extraido"].map(fold_text))
        .sort_values("Periodico_extraido")
    )
    cache = build_journal_cache(unique_journals)

    df["JCR_WoS"] = "Não"
    df["SJR_Scopus"] = "Não"
    df["PubMed_APCN"] = "Não"
    df["SciELO_APCN"] = "Não"
    df["Google_Scholar"] = "Não"
    df["Fonte_indexacao"] = ""
    df["Link_fonte"] = ""
    df["Data_consulta"] = ""
    df["Observacao_indexacao"] = ""

    for idx, row in df.iterrows():
        tipo_apcn = row["Tipo_APCN"]
        obs: list[str] = []
        if tipo_apcn == "Livro":
            df.at[idx, "Categoria_livro_capitulo"] = "Livro B"
            df.at[idx, "Pontuacao_APCN"] = 90
            df.at[idx, "Data_consulta"] = TODAY_STR
            obs.append("Livro classificado conservadoramente como Livro B conforme instrução.")
        elif tipo_apcn == "Capítulo":
            df.at[idx, "Categoria_livro_capitulo"] = "Capítulo B"
            df.at[idx, "Pontuacao_APCN"] = 45
            df.at[idx, "Data_consulta"] = TODAY_STR
            obs.append("Capítulo classificado conservadoramente como Capítulo B conforme instrução.")
        elif tipo_apcn == "Artigo em periódico":
            journal = row["Periodico_extraido"]
            if journal:
                lookup = cache.get(fold_text(journal), {})
                df.at[idx, "JCR_WoS"] = lookup.get("jcr_wos", "Não")
                df.at[idx, "SJR_Scopus"] = lookup.get("sjr_scopus", "Não")
                df.at[idx, "PubMed_APCN"] = lookup.get("pubmed", "Não")
                df.at[idx, "SciELO_APCN"] = lookup.get("scielo", "Não")
                df.at[idx, "Google_Scholar"] = lookup.get("google_scholar", MANUAL)
                df.at[idx, "Data_consulta"] = TODAY_STR

                sources: list[str] = []
                links: list[str] = []
                if lookup.get("sjr_scopus") == "Sim":
                    sources.append("SCImago Journal Rank")
                    links.append(lookup.get("scimago_link", ""))
                if lookup.get("pubmed") == "Sim":
                    sources.append("NLM Catalog")
                    links.append(lookup.get("nlm_link", ""))
                if not sources and lookup.get("jcr_wos") == MANUAL:
                    obs.append("Há indicação herdada de WoS/JCR em base anterior, mas sem evidência primária reconfirmada.")
                if lookup.get("observation"):
                    obs.append(lookup["observation"])
                df.at[idx, "Fonte_indexacao"] = " | ".join([src for src in sources if src])
                df.at[idx, "Link_fonte"] = " | ".join([link for link in links if link])
                if df.at[idx, "SJR_Scopus"] == "Sim" or df.at[idx, "PubMed_APCN"] == "Sim" or df.at[idx, "SciELO_APCN"] == "Sim":
                    df.at[idx, "Google_Scholar"] = "Não"
            else:
                df.at[idx, "JCR_WoS"] = MANUAL
                df.at[idx, "SJR_Scopus"] = MANUAL
                df.at[idx, "PubMed_APCN"] = MANUAL
                df.at[idx, "SciELO_APCN"] = MANUAL
                df.at[idx, "Google_Scholar"] = MANUAL
                df.at[idx, "Data_consulta"] = TODAY_STR
                obs.append("Periódico não identificado; indexação depende de revisão manual.")

            score = 0
            if df.at[idx, "JCR_WoS"] == "Sim" or df.at[idx, "SJR_Scopus"] == "Sim":
                score = 90
            elif df.at[idx, "PubMed_APCN"] == "Sim" or df.at[idx, "SciELO_APCN"] == "Sim":
                score = 60
            elif df.at[idx, "Google_Scholar"] == "Sim":
                score = 30
            df.at[idx, "Pontuacao_APCN"] = score

            if not row["ISSN_identificado"]:
                obs.append("ISSN ausente.")
            if not row["DOI_identificado"]:
                obs.append("DOI ausente.")
            if row["Possivel_inconsistencia_tipo"] == "Sim":
                obs.append("Possível inconsistência entre tipo informado e padrão do registro.")
        else:
            df.at[idx, "Pontuacao_APCN"] = 0
            df.at[idx, "Data_consulta"] = TODAY_STR
            obs.append("Tipo não classificado nos critérios APCN solicitados.")

        df.at[idx, "Observacao_indexacao"] = " ".join(obs).strip()

    return df


def build_sintese(df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for docente, group in df.groupby("Docente", dropna=False):
        sorted_group = group.sort_values("Pontuacao_APCN", ascending=False, kind="mergesort")
        top5 = sorted_group.head(5)
        jcr_sjr_top5 = top5[
            (top5["Tipo_APCN"] == "Artigo em periódico")
            & ((top5["JCR_WoS"] == "Sim") | (top5["SJR_Scopus"] == "Sim"))
        ]
        top5_points = pd.to_numeric(top5["Pontuacao_APCN"], errors="coerce").fillna(0).sum()
        rows.append(
            {
                "Docente": docente,
                "Total de produtos informados": int(len(group)),
                "Total de artigos": int((group["Tipo_APCN"] == "Artigo em periódico").sum()),
                "Total de livros": int((group["Tipo_APCN"] == "Livro").sum()),
                "Total de capítulos": int((group["Tipo_APCN"] == "Capítulo").sum()),
                "Número de produtos com 90 pontos": int((group["Pontuacao_APCN"] == 90).sum()),
                "Número de produtos com 60 pontos": int((group["Pontuacao_APCN"] == 60).sum()),
                "Número de produtos com 30 pontos": int((group["Pontuacao_APCN"] == 30).sum()),
                "Número de produtos com 0 ponto": int((group["Pontuacao_APCN"] == 0).sum()),
                "Pontuação total considerando todos os produtos": int(
                    pd.to_numeric(group["Pontuacao_APCN"], errors="coerce").fillna(0).sum()
                ),
                "Pontuação dos 5 melhores produtos APCN": int(top5_points),
                "Número de produtos JCR_WoS ou SJR_Scopus entre os 5 melhores": int(len(jcr_sjr_top5)),
                "Atende_330_pontos_Doutorado": "Sim" if top5_points >= 330 else "Não",
                "Atende_2_produtos_JCR_SJR_Doutorado": "Sim" if len(jcr_sjr_top5) >= 2 else "Não",
            }
        )
    return pd.DataFrame(rows).sort_values("Docente").reset_index(drop=True)


def build_pendencias(df: pd.DataFrame) -> pd.DataFrame:
    def has_pending(row: pd.Series) -> bool:
        manual_fields = ["JCR_WoS", "SJR_Scopus", "PubMed_APCN", "SciELO_APCN", "Google_Scholar"]
        if any(row.get(field) == MANUAL for field in manual_fields):
            return True
        if row.get("Tipo_APCN") == "Artigo em periódico" and not row.get("Periodico_extraido"):
            return True
        if row.get("Tipo_APCN") == "Artigo em periódico" and not row.get("ISSN_identificado"):
            return True
        if row.get("Tipo_APCN") == "Artigo em periódico" and not row.get("DOI_identificado"):
            return True
        if row.get("Possivel_inconsistencia_tipo") == "Sim":
            return True
        if not row.get("Fonte_indexacao") and row.get("Tipo_APCN") == "Artigo em periódico":
            return True
        return False

    pend = df[df.apply(has_pending, axis=1)].copy()
    reasons = []
    for _, row in pend.iterrows():
        row_reasons = []
        if row["Tipo_APCN"] == "Artigo em periódico" and not row["Periodico_extraido"]:
            row_reasons.append("Periódico não identificado")
        if row["Tipo_APCN"] == "Artigo em periódico" and not row["ISSN_identificado"]:
            row_reasons.append("ISSN ausente")
        if row["Tipo_APCN"] == "Artigo em periódico" and not row["DOI_identificado"]:
            row_reasons.append("DOI ausente")
        if row["Possivel_inconsistencia_tipo"] == "Sim":
            row_reasons.append("Produto com possível inconsistência de tipo")
        if not row["Fonte_indexacao"] and row["Tipo_APCN"] == "Artigo em periódico":
            row_reasons.append("Fonte não confirmada")
        if any(row.get(field) == MANUAL for field in ["JCR_WoS", "SJR_Scopus", "PubMed_APCN", "SciELO_APCN", "Google_Scholar"]):
            row_reasons.append("Indexação ambígua")
        reasons.append(" | ".join(row_reasons))
    pend["Motivo_pendencia"] = reasons
    return pend


def autosize_sheet(ws) -> None:
    for col_cells in ws.columns:
        max_len = 0
        col_index = col_cells[0].column
        for cell in col_cells[:200]:
            if cell.value is None:
                continue
            max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col_index)].width = min(max(max_len + 2, 10), 42)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def add_table(ws, table_name: str) -> None:
    if ws.max_row < 2 or ws.max_column < 1:
        return
    tab = Table(displayName=table_name, ref=ws.dimensions)
    style = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    tab.tableStyleInfo = style
    ws.add_table(tab)


def style_header(ws) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font


def main() -> None:
    source_df = pd.read_excel(INPUT_XLSX)
    classified = classify_articles(source_df.copy())
    output_columns = list(source_df.columns)
    derived_columns = [
        "Tipo_APCN",
        "Periodico_extraido",
        "ISSN_identificado",
        "DOI_identificado",
        "JCR_WoS",
        "SJR_Scopus",
        "PubMed_APCN",
        "SciELO_APCN",
        "Google_Scholar",
        "Categoria_livro_capitulo",
        "Pontuacao_APCN",
        "Fonte_indexacao",
        "Link_fonte",
        "Data_consulta",
        "Observacao_indexacao",
    ]
    final_df = classified[output_columns + derived_columns].copy()
    sintese = build_sintese(classified)
    pendencias = build_pendencias(classified)

    with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as writer:
        final_df.to_excel(writer, index=False, sheet_name="Producao_classificada")
        sintese.to_excel(writer, index=False, sheet_name="Sintese_docente")
        pendencias.to_excel(writer, index=False, sheet_name="Pendencias")

    wb = load_workbook(OUTPUT_XLSX)
    for ws in wb.worksheets:
        style_header(ws)
        autosize_sheet(ws)
    add_table(wb["Producao_classificada"], "TabelaProducaoClassificada")
    add_table(wb["Sintese_docente"], "TabelaSinteseDocente")
    add_table(wb["Pendencias"], "TabelaPendencias")
    wb.save(OUTPUT_XLSX)

    summary = {
        "output": str(OUTPUT_XLSX),
        "rows_producao": int(len(final_df)),
        "rows_pendencias": int(len(pendencias)),
        "docentes": int(sintese["Docente"].nunique()),
        "scores": final_df["Pontuacao_APCN"].value_counts(dropna=False).to_dict(),
        "sjr_sim": int((final_df["SJR_Scopus"] == "Sim").sum()),
        "pubmed_sim": int((final_df["PubMed_APCN"] == "Sim").sum()),
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
